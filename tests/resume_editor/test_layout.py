from copy import deepcopy

import pytest
from docx import Document
from fastapi.testclient import TestClient
from pypdf import PdfReader

from jobfindsme.desktop_api import create_app
from jobfindsme.resume_editor import ResumeEditorService
from jobfindsme.resume_editor.document import reconcile_document, validate_document
from tests.resume_editor.test_service import _confirmed_resume

CONTENT = {
    "basic_information": ["Fictional Candidate"],
    "education": ["Example University"],
    "experience": ["Example Engineering"],
    "projects": ["Example Project"],
    "skills": ["Python FastAPI"],
}


def layout_document(template):
    document = reconcile_document(CONTENT)
    document["attrs"] = {"template": template}
    document["content"][0]["content"][0]["content"][0]["marks"] = [
        {"type": "bold"},
        {"type": "textStyle", "attrs": {"color": "#1769aa", "fontSize": 20}},
    ]
    return document


@pytest.mark.parametrize("template", ["campus", "compact", "professional"])
def test_default_exports_follow_saved_layout_and_retain_typography(tmp_path, template):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    saved = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=CONTENT,
        document=layout_document(template),
    )
    service = ResumeEditorService(db)
    assert service.get_version(
        workspace_id=workspace.workspace_id, version_id=saved.version_id
    ).document["attrs"] == {"template": template}
    expected = (
        [
            "Fictional Candidate",
            "Example University",
            "Example Project",
            "Example Engineering",
            "Python FastAPI",
        ]
        if template == "campus"
        else [
            "Fictional Candidate",
            "Example Engineering",
            "Example Project",
            "Python FastAPI",
            "Example University",
        ]
    )
    for format in ("pdf", "docx"):
        output = tmp_path / f"{template}.{format}"
        exported = service.export(
            workspace_id=workspace.workspace_id,
            version_id=saved.version_id,
            destination=output,
            format=format,
        )
        assert exported.template == template
        text = (
            "\n".join(p.text for p in Document(output).paragraphs)
            if format == "docx"
            else "\n".join(p.extract_text() for p in PdfReader(output).pages)
        )
        positions = [text.index(value) for value in expected]
        assert positions == sorted(positions)
        assert "•" not in text
        if format == "docx":
            word = Document(output)
            name = next(
                r for p in word.paragraphs for r in p.runs if r.text == expected[0]
            )
            assert name.bold and name.font.size.pt == 20
            assert str(name.font.color.rgb) == "1769AA"
            assert all(p.style.name != "List Bullet" for p in word.paragraphs)
            assert word.sections[0].left_margin.mm == pytest.approx(
                16.5 if template == "compact" else 20, abs=0.1
            )
    assert (
        profiles.list_versions(workspace_id=workspace.workspace_id)[0].document[
            "attrs"
        ]["template"]
        == template
    )


def test_layout_is_versioned_inherited_by_edits_and_restored_without_rewriting_history(
    tmp_path,
):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    first = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=CONTENT,
        document=layout_document("campus"),
    )
    compact = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=first.version_id,
        content=CONTENT,
        document=layout_document("compact"),
    )
    edited = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=compact.version_id,
        content={**CONTENT, "skills": ["Python FastAPI Redis"]},
    )
    assert edited.document["attrs"]["template"] == "compact"
    restored = service.restore_version(
        workspace_id=workspace.workspace_id, version_id=first.version_id
    )
    assert restored.document["attrs"]["template"] == "campus"
    assert restored.content == first.content
    assert (
        service.get_version(
            workspace_id=workspace.workspace_id, version_id=compact.version_id
        ).document
        == compact.document
    )
    for version, expected in (
        (first, "campus"),
        (compact, "compact"),
        (restored, "campus"),
    ):
        assert (
            service.export(
                workspace_id=workspace.workspace_id,
                version_id=version.version_id,
                destination=tmp_path / f"{version.version_id}.pdf",
                format="pdf",
            ).template
            == expected
        )
    assert (
        service.get_version(
            workspace_id=workspace.workspace_id, version_id=base.version_id
        ).document
        == base.document
    )


def test_document_accepts_only_known_layouts_and_preserves_them_during_reconciliation():
    original = layout_document("campus")
    assert validate_document(original)["attrs"] == {"template": "campus"}
    assert reconcile_document(CONTENT, original)["attrs"] == {"template": "campus"}
    bad = deepcopy(original)
    bad["attrs"]["template"] = "external-style"
    with pytest.raises(ValueError, match="invalid resume template"):
        validate_document(bad)


def test_desktop_api_exports_selected_version_layout_when_caller_omits_template(
    tmp_path,
):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    saved = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=CONTENT,
        document=layout_document("campus"),
    )
    headers = {"Authorization": "Bearer test"}
    with TestClient(create_app(token="test", database_path=db.path)) as client:
        for format in ("pdf", "docx"):
            destination = tmp_path / f"api.{format}"
            response = client.post(
                f"/v1/resume-versions/{saved.version_id}/export",
                headers=headers,
                json={
                    "workspace_id": workspace.workspace_id,
                    "format": format,
                    "destination": str(destination),
                },
            )
            assert response.status_code == 200
            assert response.json()["template"] == "campus"
            assert destination.exists()
        legacy = client.post(
            f"/v1/resume-versions/{base.version_id}/export",
            headers=headers,
            json={
                "workspace_id": workspace.workspace_id,
                "format": "pdf",
                "destination": str(tmp_path / "legacy.pdf"),
            },
        )
        assert legacy.status_code == 200 and legacy.json()["template"] == "professional"
