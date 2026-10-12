from jobfindsme.profiles.models import FactType
from jobfindsme.profiles.parser import DeterministicResumeParser


def test_parser_returns_deduplicated_facts_with_exact_evidence() -> None:
    text = """# Skills
Python, FastAPI, Python, RAG

# Projects
- jobfindsme uses MCP for local job search
"""

    facts = DeterministicResumeParser().parse(text)

    assert [fact.value for fact in facts].count("Python") == 1
    assert any(fact.fact_type is FactType.PROJECT for fact in facts)
    for fact in facts:
        assert text[fact.evidence_start : fact.evidence_end] == (fact.evidence_snippet)


def test_parser_v3_groups_wrapped_resume_sections_without_line_noise() -> None:
    text = """独立实践与工作经历
2025.01-2025.06 jobfindsme
负责 Agent 与 MCP 工具设计，
并使用 Pydantic、pytest 和 RRF 完成质量门禁。
教育背景
2020.09-2024.06 示例大学
计算机科学与技术 本科
"""

    facts = DeterministicResumeParser().parse(text)

    experience = [fact for fact in facts if fact.fact_type is FactType.EXPERIENCE]
    education = [fact for fact in facts if fact.fact_type is FactType.EDUCATION]
    skills = {fact.value for fact in facts if fact.fact_type is FactType.SKILL}
    assert len(experience) == 1
    assert "负责 Agent 与 MCP 工具设计" in experience[0].value
    assert len(education) == 1
    assert {"Pydantic", "pytest", "RRF"} <= skills
    assert all(
        text[fact.evidence_start : fact.evidence_end] == fact.evidence_snippet
        for fact in facts
    )


def test_two_distinct_education_entries_remain_separate_facts() -> None:
    text = """教育经历
2020.09-2024.06 示例大学 计算机科学 本科
2024.09-2026.06 示例研究院 人工智能 硕士
"""
    education = [
        fact
        for fact in DeterministicResumeParser().parse(text)
        if fact.fact_type is FactType.EDUCATION
    ]
    assert len(education) == 2
    assert education[0].value.startswith("2020.09-2024.06")
    assert education[1].value.startswith("2024.09-2026.06")
    assert "示例大学" in education[0].value
    assert "示例研究院" in education[1].value
    assert all(
        text[fact.evidence_start : fact.evidence_end] == fact.evidence_snippet
        for fact in education
    )


def test_inline_section_labels_yield_reviewable_facts_without_private_basics():
    from jobfindsme.profiles.models import FactType

    text = (
        "姓名：测试用户\n求职意向：AI 应用工程师\n"
        "教育经历：示例大学 计算机科学 本科\n"
        "工作经历：示例公司 2022-2025 开发企业知识库应用\n"
        "项目经历：使用 Python 和检索增强技术实现问答系统，负责接口与评估。\n"
        "技能：Python、TypeScript、SQL、RAG\n"
    )
    facts = DeterministicResumeParser().parse(text)
    kinds = {item.fact_type for item in facts}
    assert {FactType.EDUCATION, FactType.EXPERIENCE, FactType.PROJECT} <= kinds
    assert all("测试用户" not in item.value for item in facts)
    assert all(
        "技能：" not in item.value
        for item in facts
        if item.fact_type is FactType.PROJECT
    )


def test_full_resume_preserves_local_basics_skill_descriptions_and_short_entries():
    text = (
        "个人简历\n姓名：虚构候选人\n电话：13800138000\nqa@example.com\n"
        "**工作经历**\n示例公司\n负责内部工具维护。\n"
        "教 育 背 景\n示例大学 本科\n"
        "专业技能\n熟悉 Python，能编写单元测试和排查服务问题。\n"
    )
    parser = DeterministicResumeParser()
    facts = parser.parse(text, include_basic_information=True)
    assert [f.value for f in facts if f.fact_type is FactType.BASIC_INFORMATION] == [
        "姓名：虚构候选人",
        "电话：13800138000",
        "qa@example.com",
    ]
    assert any(
        f.fact_type is FactType.SKILL_DETAIL and "单元测试" in f.value for f in facts
    )
    assert any(
        f.fact_type is FactType.EDUCATION and f.value == "示例大学 本科" for f in facts
    )
    assert any(
        f.fact_type is FactType.EXPERIENCE and "示例公司" in f.value for f in facts
    )
    assert all(
        text[f.evidence_start : f.evidence_end] == f.evidence_snippet for f in facts
    )
    assert all(
        f.fact_type is not FactType.BASIC_INFORMATION for f in parser.parse(text)
    )


def test_docx_preserves_soft_line_breaks_and_tabs():
    import io

    from docx import Document

    from jobfindsme.profiles.parser import ResumeTextExtractor

    document = Document()
    run = document.add_paragraph().add_run("工作经历")
    run.add_break()
    run.add_text("虚构公司")
    run.add_tab()
    run.add_text("后端开发，维护 Python 服务。")
    output = io.BytesIO()
    document.save(output)
    extracted = ResumeTextExtractor().extract(
        file_name="synthetic.docx", content=output.getvalue()
    )
    facts = DeterministicResumeParser().parse(extracted.text)
    assert any(
        f.fact_type is FactType.EXPERIENCE and "虚构公司" in f.value for f in facts
    )
    assert "\t" in extracted.text


def test_numbered_headings_do_not_swallow_projects_into_skills():
    text = (
        "01 专业技能\nPython、SQL\n"
        "02 项目经历\n虚构工具 2024.03-2025.01\n"
        "负责接口维护。\n编写回归测试。\n"
        "03 开源贡献\n参与示例文档修订。\n"
        "（四）教 育 背 景\n示例大学 本科\n"
    )
    facts = DeterministicResumeParser().parse(text)
    skills = [f.value for f in facts if f.fact_type is FactType.SKILL_DETAIL]
    assert skills == ["Python、SQL"]
    projects = [f.value for f in facts if f.fact_type is FactType.PROJECT]
    assert "负责接口维护。\n编写回归测试。" in projects[0]
    assert any("示例文档" in value for value in projects)
    assert any(f.fact_type is FactType.EDUCATION for f in facts)
    assert all(
        text[f.evidence_start : f.evidence_end] == f.evidence_snippet for f in facts
    )


def test_numbered_achievements_and_dates_are_not_section_headings():
    text = "项目经历\n示例项目 2024.01-2025.01\n01 完成接口整理\n02 验证接口返回"
    projects = [
        f.value
        for f in DeterministicResumeParser().parse(text)
        if f.fact_type is FactType.PROJECT
    ]
    assert len(projects) == 1
    assert "01 完成接口整理" in projects[0]
    assert "02 验证接口返回" in projects[0]


def test_nested_technology_and_description_do_not_steal_project_body():
    text = (
        "项目经历\n虚构检索项目 2024.01-2024.12\n"
        "项目介绍：供内部使用的检索工具。\n技术栈：Python、Redis\n"
        "负责索引更新。\n在2024.03-2024.06期间补充测试。\n"
        "虚构助手项目 2025.01-至今\n技术栈：TypeScript\n负责交互设计。\n"
        "工作经历\n虚构公司 2023.01-2025.01\n技术栈：SQL\n负责报表。\n"
        "专业技能\n熟悉 Python。\n能够独立维护服务。\n"
    )
    facts = DeterministicResumeParser().parse(text)
    projects = [f for f in facts if f.fact_type is FactType.PROJECT]
    assert len(projects) == 2
    assert "技术栈：Python、Redis" in projects[0].value
    assert "项目介绍：供内部使用" in projects[0].value
    assert "补充测试" in projects[0].value
    assert "负责交互设计" in projects[1].value
    assert [f.value for f in facts if f.fact_type is FactType.SKILL_DETAIL] == [
        "熟悉 Python。\n能够独立维护服务。"
    ]
    assert any(
        "负责报表" in f.value for f in facts if f.fact_type is FactType.EXPERIENCE
    )
    assert all(
        text[f.evidence_start : f.evidence_end] == f.evidence_snippet for f in facts
    )


def test_unindexed_sections_and_repeated_authored_lines_are_not_discarded():
    text = (
        "自我评价\n重视协作。\n证书：示例证书\n专业技能\nPython\n"
        "项目经历\n示例项目\n负责测试\n工作经历\n负责测试\n"
    )
    parser = DeterministicResumeParser()
    facts = parser.parse(text, include_basic_information=True)
    basics = [f.value for f in facts if f.fact_type is FactType.BASIC_INFORMATION]
    assert basics == ["自我评价", "重视协作。", "证书：示例证书"]
    assert all(
        f.fact_type is not FactType.BASIC_INFORMATION for f in parser.parse(text)
    )
    repeated = parser.parse(
        "基本信息\n同一行\n同一行\n", include_basic_information=True
    )
    assert len(repeated) == 2


def test_imported_pdf_bullets_survive_section_parsing():
    from jobfindsme.profiles.text import normalize_resume_text

    text = normalize_resume_text(
        "工作经历\n示例公司 2024.01 - 2025.01\n"
        "\uf0b7 完成 Python 服务\n\uf0b7 优化检索性能"
    )
    facts = DeterministicResumeParser().parse(text)
    experience = next(f for f in facts if f.fact_type is FactType.EXPERIENCE)
    assert experience.value == (
        "示例公司 2024.01 - 2025.01\n• 完成 Python 服务\n• 优化检索性能"
    )
    assert experience.evidence_snippet == experience.value
