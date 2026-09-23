from __future__ import annotations

import json
import os
from dataclasses import dataclass
from datetime import UTC, datetime
from pathlib import Path
from typing import Literal
from uuid import uuid4

from jobfindsme.profiles.models import ResumeVersion
from jobfindsme.storage import Database

ResumeTemplate = Literal["classic", "compact"]
ResumeFormat = Literal["pdf", "docx", "md"]
SECTION_ORDER = (
    "basic_information",
    "education",
    "experience",
    "projects",
    "skills",
)
SECTION_TITLES = {
    "basic_information": "基本信息",
    "education": "教育经历",
    "experience": "工作经历",
    "projects": "项目经历",
    "skills": "专业技能",
}


class ResumeEditorError(ValueError):
    pass


@dataclass(frozen=True)
class ResumeExport:
    path: Path
    format: ResumeFormat
    template: ResumeTemplate
    version_id: str


class ResumeEditorService:
    def __init__(self, database: Database) -> None:
        self.database = database

    def list_versions(self, *, workspace_id: str) -> list[ResumeVersion]:
        self.database.migrate()
        with self.database.connect() as connection:
            rows = connection.execute(
                """SELECT * FROM resume_versions WHERE workspace_id = ?
                AND hidden_at IS NULL
                ORDER BY version_number DESC""",
                (workspace_id,),
            ).fetchall()
        return [_version_from_row(row) for row in rows]

    def hide_version(self, *, workspace_id: str, version_id: str) -> None:
        """Hide one historical version while retaining its immutable references."""
        self.database.migrate()
        with self.database.connect() as connection:
            connection.execute("BEGIN IMMEDIATE")
            row = connection.execute(
                "SELECT is_current, hidden_at FROM resume_versions "
                "WHERE workspace_id=? AND version_id=?",
                (workspace_id, version_id),
            ).fetchone()
            if row is None or row["hidden_at"] is not None:
                raise ResumeEditorError("resume version not found")
            if row["is_current"]:
                raise ResumeEditorError("current resume version cannot be hidden")
            for table in (
                "desktop_search_runs", "desktop_scheduled_tasks", "research_reports"
            ):
                if connection.execute(
                    f"SELECT 1 FROM {table} WHERE workspace_id=? "
                    "AND resume_version_id=? LIMIT 1",
                    (workspace_id, version_id),
                ).fetchone():
                    raise ResumeEditorError(
                        "referenced resume version cannot be hidden"
                    )
            connection.execute(
                "UPDATE resume_versions SET hidden_at=? "
                "WHERE workspace_id=? AND version_id=? AND is_current=0",
                (datetime.now(UTC).isoformat(), workspace_id, version_id),
            )

    def get_version(self, *, workspace_id: str, version_id: str) -> ResumeVersion:
        self.database.migrate()
        with self.database.connect() as connection:
            row = connection.execute(
                """SELECT * FROM resume_versions
                WHERE workspace_id = ? AND version_id = ?""",
                (workspace_id, version_id),
            ).fetchone()
        if row is None:
            raise ResumeEditorError("resume version not found")
        return _version_from_row(row)

    def save_edit(
        self,
        *,
        workspace_id: str,
        base_version_id: str,
        content: dict[str, list[str] | tuple[str, ...]],
    ) -> ResumeVersion:
        normalized = _normalize_content(content)
        return self._create_from_base(
            workspace_id=workspace_id,
            base_version_id=base_version_id,
            content=normalized,
        )

    def restore_version(
        self,
        *,
        workspace_id: str,
        version_id: str,
    ) -> ResumeVersion:
        target = self.get_version(workspace_id=workspace_id, version_id=version_id)
        if not any(
            version.version_id == target.version_id
            for version in self.list_versions(workspace_id=workspace_id)
        ):
            raise ResumeEditorError("hidden resume version cannot be restored")
        with self.database.connect() as connection:
            current = connection.execute(
                """SELECT version_id FROM resume_versions
                WHERE workspace_id = ? AND is_current = 1""",
                (workspace_id,),
            ).fetchone()
        if current is None:
            raise ResumeEditorError("current resume version not found")
        return self._create_from_base(
            workspace_id=workspace_id,
            base_version_id=current["version_id"],
            content=target.content,
        )

    def _create_from_base(
        self,
        *,
        workspace_id: str,
        base_version_id: str,
        content: dict[str, tuple[str, ...]],
    ) -> ResumeVersion:
        with self.database.connect() as connection:
            return self._create_from_base_in_transaction(
                connection=connection,
                workspace_id=workspace_id,
                base_version_id=base_version_id,
                content=content,
            )

    def _create_from_base_in_transaction(
        self,
        *,
        connection,
        workspace_id: str,
        base_version_id: str,
        content: dict[str, tuple[str, ...]],
    ) -> ResumeVersion:
        """Create a version on a caller-owned transaction.

        Prompt-session saves use this so advancing the current version and
        closing the edit session commit or roll back together.
        """
        now = datetime.now(UTC)
        version_id = f"resume_version_{uuid4().hex}"
        base = connection.execute(
            """SELECT * FROM resume_versions
            WHERE workspace_id = ? AND version_id = ?""",
            (workspace_id, base_version_id),
        ).fetchone()
        if base is None:
            raise ResumeEditorError("base resume version not found")
        current = connection.execute(
            """SELECT version_id FROM resume_versions
            WHERE workspace_id = ? AND is_current = 1""",
            (workspace_id,),
        ).fetchone()
        if current is None or current["version_id"] != base_version_id:
            raise ResumeEditorError(
                "resume version conflict: reload the current version before saving"
            )
        next_number = connection.execute(
            """SELECT COALESCE(MAX(version_number), 0) + 1
            FROM resume_versions WHERE workspace_id = ?""",
            (workspace_id,),
        ).fetchone()[0]
        connection.execute(
            "UPDATE resume_versions SET is_current = 0 WHERE workspace_id = ?",
            (workspace_id,),
        )
        connection.execute(
            """INSERT INTO resume_versions (
                version_id, workspace_id, profile_id, source_document_id,
                parent_version_id, version_number, content_json,
                is_current, created_at
            ) VALUES (?, ?, ?, ?, ?, ?, ?, 1, ?)""",
            (
                version_id,
                workspace_id,
                base["profile_id"],
                base["source_document_id"],
                base_version_id,
                next_number,
                json.dumps(content, ensure_ascii=False),
                now.isoformat(),
            ),
        )
        created = connection.execute(
            "SELECT * FROM resume_versions WHERE version_id = ?",
            (version_id,),
        ).fetchone()
        return _version_from_row(created)

    def export(
        self,
        *,
        workspace_id: str,
        version_id: str,
        destination: str | Path,
        format: ResumeFormat,
        template: ResumeTemplate,
    ) -> ResumeExport:
        if format not in {"pdf", "docx", "md"}:
            raise ResumeEditorError("unsupported export format")
        if template not in {"classic", "compact"}:
            raise ResumeEditorError("unsupported resume template")
        version = self.get_version(workspace_id=workspace_id, version_id=version_id)
        path = Path(destination).expanduser()
        if path.suffix.casefold() != f".{format}":
            raise ResumeEditorError(f"export path must end with .{format}")
        path.parent.mkdir(parents=True, exist_ok=True)
        if format == "md":
            path.write_text(_markdown(version.content, template), encoding="utf-8")
        elif format == "docx":
            _write_docx(path, version.content, template)
        else:
            _write_pdf(path, version.content, template)
        os.chmod(path, 0o600)
        return ResumeExport(path, format, template, version.version_id)


def _normalize_content(
    content: dict[str, list[str] | tuple[str, ...]],
) -> dict[str, tuple[str, ...]]:
    unknown = set(content) - set(SECTION_ORDER)
    if unknown:
        raise ResumeEditorError(f"unknown resume sections: {sorted(unknown)}")
    normalized: dict[str, tuple[str, ...]] = {}
    for section in SECTION_ORDER:
        values = content.get(section, ())
        if not isinstance(values, (list, tuple)):
            raise ResumeEditorError(f"resume section {section} must be a list")
        # Preserve the editor's line breaks and spacing exactly across versions.
        cleaned = tuple(str(value) for value in values)
        normalized[section] = cleaned
    if not any(value.strip() for values in normalized.values() for value in values):
        raise ResumeEditorError("resume content must not be empty")
    return normalized


def _markdown(content: dict[str, tuple[str, ...]], template: ResumeTemplate) -> str:
    lines = ["# 简历", ""]
    for section in SECTION_ORDER:
        values = content.get(section, ())
        if not values:
            continue
        lines.extend([f"## {SECTION_TITLES[section]}", ""])
        marker = "- " if template == "classic" else ""
        for value in values:
            lines.extend([f"{marker}{value}", ""])
    return "\n".join(lines).rstrip() + "\n"


def _write_docx(
    path: Path,
    content: dict[str, tuple[str, ...]],
    template: ResumeTemplate,
) -> None:
    try:
        from docx import Document
        from docx.enum.text import WD_LINE_SPACING
        from docx.oxml.ns import qn
        from docx.shared import Cm, Pt, RGBColor
    except ImportError as error:
        raise ResumeEditorError(
            "DOCX export requires the python-docx dependency"
        ) from error
    document = Document()
    section = document.sections[0]
    section.top_margin = Cm(1.45 if template == "compact" else 1.8)
    section.bottom_margin = section.top_margin
    section.left_margin = Cm(1.65 if template == "compact" else 2.0)
    section.right_margin = section.left_margin
    normal = document.styles["Normal"]
    normal.font.name = "Arial"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "PingFang SC")
    normal.font.size = Pt(9.5 if template == "compact" else 10.5)
    title = document.add_heading("简历", level=0)
    title.alignment = 1
    for name in SECTION_ORDER:
        values = content.get(name, ())
        if not values:
            continue
        heading = document.add_heading(SECTION_TITLES[name], level=1)
        heading.runs[0].font.color.rgb = RGBColor(48, 48, 46)
        for value in values:
            paragraph = document.add_paragraph(style=None)
            if template == "classic":
                paragraph.style = document.styles["List Bullet"]
            paragraph.paragraph_format.space_after = Pt(
                2 if template == "compact" else 5
            )
            paragraph.paragraph_format.line_spacing_rule = WD_LINE_SPACING.SINGLE
            paragraph.add_run(value)
    document.save(path)


def _write_pdf(
    path: Path,
    content: dict[str, tuple[str, ...]],
    template: ResumeTemplate,
) -> None:
    try:
        from reportlab.lib import colors
        from reportlab.lib.enums import TA_CENTER
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
        from reportlab.lib.units import mm
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
    except ImportError as error:
        raise ResumeEditorError(
            "PDF export requires the reportlab dependency"
        ) from error
    font_name = _register_cjk_font(pdfmetrics, TTFont)
    margin = (15 if template == "compact" else 19) * mm
    document = SimpleDocTemplate(
        str(path),
        pagesize=A4,
        leftMargin=margin,
        rightMargin=margin,
        topMargin=margin,
        bottomMargin=margin,
        title="简历",
        author="JobFindsMe",
    )
    styles = getSampleStyleSheet()
    title = ParagraphStyle(
        "ResumeTitle",
        parent=styles["Title"],
        fontName=font_name,
        fontSize=18,
        leading=23,
        alignment=TA_CENTER,
        textColor=colors.HexColor("#30302e"),
        spaceAfter=8 * mm,
    )
    heading = ParagraphStyle(
        "ResumeHeading",
        parent=styles["Heading2"],
        fontName=font_name,
        fontSize=11.5,
        leading=15,
        textColor=colors.HexColor("#30302e"),
        spaceBefore=4 * mm,
        spaceAfter=2 * mm,
        borderWidth=0,
        borderPadding=0,
    )
    body = ParagraphStyle(
        "ResumeBody",
        parent=styles["BodyText"],
        fontName=font_name,
        fontSize=8.8 if template == "compact" else 9.8,
        leading=13 if template == "compact" else 15,
        textColor=colors.HexColor("#3f3f3c"),
        leftIndent=0 if template == "compact" else 3 * mm,
        spaceAfter=1.5 * mm if template == "compact" else 2.5 * mm,
        wordWrap="CJK",
    )
    story = [Paragraph("简历", title)]
    for name in SECTION_ORDER:
        values = content.get(name, ())
        if not values:
            continue
        story.extend([Paragraph(SECTION_TITLES[name], heading), Spacer(1, 0.5 * mm)])
        for value in values:
            prefix = "• " if template == "classic" else ""
            story.append(Paragraph(prefix + _xml_escape(value), body))
    document.build(story)


def _register_cjk_font(pdfmetrics, ttfont) -> str:
    candidates = (
        Path("/System/Library/Fonts/STHeiti Light.ttc"),
        Path("/System/Library/Fonts/Supplemental/Songti.ttc"),
        Path("/usr/share/fonts/opentype/noto/NotoSansCJK-Regular.ttc"),
    )
    for candidate in candidates:
        if candidate.exists():
            try:
                pdfmetrics.registerFont(
                    ttfont("JobFindsMeCJK", str(candidate), subfontIndex=0)
                )
                return "JobFindsMeCJK"
            except Exception:
                continue
    raise ResumeEditorError(
        "PDF export requires an installed Chinese font (PingFang/Heiti/Noto CJK)"
    )


def _xml_escape(value: str) -> str:
    return value.replace("&", "&amp;").replace("<", "&lt;").replace(">", "&gt;")


def _version_from_row(row) -> ResumeVersion:
    content = json.loads(row["content_json"])
    return ResumeVersion(
        version_id=row["version_id"],
        workspace_id=row["workspace_id"],
        profile_id=row["profile_id"],
        source_document_id=row["source_document_id"],
        parent_version_id=row["parent_version_id"],
        version_number=row["version_number"],
        content={key: tuple(values) for key, values in content.items()},
        is_current=bool(row["is_current"]),
        created_at=datetime.fromisoformat(row["created_at"]),
    )
