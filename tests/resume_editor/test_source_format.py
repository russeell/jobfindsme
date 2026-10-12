from copy import deepcopy
from pathlib import Path

import pytest
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTChar
from reportlab.pdfgen import canvas

from jobfindsme.profiles.models import ResumeImportMode
from jobfindsme.resume_editor import ResumeEditorService
from jobfindsme.resume_editor.document import (
    document_content,
    reconcile_document,
    validate_document,
)
from jobfindsme.resume_editor.source_format import import_document
from tests.resume_editor.test_service import _confirmed_resume


def _texts(document):
    return [
        n
        for s in document["content"]
        for b in s["content"]
        for n in b.get("content", [])
        if n["type"] == "text"
    ]


def _style(node):
    return next(m["attrs"] for m in node["marks"] if m["type"] == "textStyle")


def _colored_pdf(path):
    c = canvas.Canvas(str(path))
    c.setFillColorRGB(0.1, 0.2, 0.3)
    c.setFont("Helvetica-Bold", 24)
    c.drawString(50, 780, "Example Candidate")
    c.setFillColorRGB(0, 0.5, 0.4)
    c.setFont("Helvetica", 12)
    c.drawString(50, 750, "AI Engineer")
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, 710, "Skills")
    c.setFont("Helvetica", 10)
    c.setFillColorRGB(0.1, 0.3, 0.8)
    c.drawString(50, 690, "Python FastAPI")
    c.save()


def test_existing_original_gets_color_size_and_bold_without_history_rewrite(tmp_path):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    path = tmp_path / "source.pdf"
    _colored_pdf(path)
    raw = path.read_bytes()
    profiles.import_resume(
        workspace_id=workspace.workspace_id,
        source_path=path,
        mode=ResumeImportMode.MANAGED,
        index_original=True,
    )
    with db.connect() as conn:
        before = tuple(
            conn.execute(
                "SELECT content_json,document_json FROM resume_versions "
                "WHERE is_current=1"
            ).fetchone()
        )
    original = profiles.current_version(workspace_id=workspace.workspace_id)
    editor = ResumeEditorService(db)
    loaded = editor.get_version(
        workspace_id=workspace.workspace_id, version_id=original.version_id
    )
    assert loaded.document == original.document
    assert document_content(loaded.document) == loaded.content
    texts = _texts(loaded.document)
    name = next(n for n in texts if n["text"] == "Example Candidate")
    assert {"type": "bold"} in name["marks"]
    assert _style(name) == {"color": "#1a334c", "fontSize": 24.0}
    skill = next(n for n in texts if n["text"] == "Python FastAPI")
    assert _style(skill)["color"] == "#1a4ccc"
    assert loaded.document["content"][0]["attrs"]["headingHidden"]
    changed = editor.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=loaded.version_id,
        content={**loaded.content, "skills": ["Python FastAPI SQLite"]},
    )
    assert _style(_texts(changed.document)[-1])["color"] == "#1a4ccc"
    for format in ("docx", "pdf"):
        output = tmp_path / f"export.{format}"
        editor.export(
            workspace_id=workspace.workspace_id,
            version_id=changed.version_id,
            format=format,
            template="professional",
            destination=output,
        )
    word = Document(tmp_path / "export.docx")
    name_run = next(
        r for p in word.paragraphs for r in p.runs if r.text == "Example Candidate"
    )
    assert name_run.bold and name_run.font.size.pt == 24
    assert str(name_run.font.color.rgb) == "1A334C"
    exported = []

    def walk(node):
        if isinstance(node, LTChar):
            exported.append(node)
        elif hasattr(node, "__iter__"):
            for child in node:
                walk(child)

    for page in extract_pages(tmp_path / "export.pdf"):
        walk(page)
    assert any(round(n.size, 1) == 24 for n in exported)
    assert any(
        n.graphicstate.ncolor
        == pytest.approx((26 / 255, 76 / 255, 204 / 255), abs=1e-6)
        for n in exported
    )
    assert path.read_bytes() == raw
    with db.connect() as conn:
        after = tuple(
            conn.execute(
                "SELECT content_json,document_json FROM resume_versions "
                "WHERE version_id=?",
                (loaded.version_id,),
            ).fetchone()
        )
    assert after == before


def test_docx_mixed_runs_inherit_style_alignment_and_export(tmp_path):
    word = Document()
    word.styles["Normal"].font.size = Pt(11)
    word.styles["Normal"].font.color.rgb = RGBColor.from_string("17324D")
    p = word.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.add_run("Candidate ").bold = True
    p.add_run("Name").italic = True
    word.add_paragraph("Skills")
    p = word.add_paragraph()
    p.add_run("Python").font.color.rgb = RGBColor.from_string("0F766E")
    p.add_run(" / SQL").underline = True
    path = tmp_path / "sample.docx"
    word.save(path)
    content = document_content(
        reconcile_document(
            {"basic_information": ["Candidate Name"], "skills": ["Python / SQL"]}
        )
    )
    rich = import_document(content, path.read_bytes(), ".docx")
    assert document_content(rich) == content
    basic = rich["content"][0]["content"][0]
    assert basic["attrs"]["textAlign"] == "center"
    assert _style(basic["content"][0])["fontSize"] == 11
    assert {"type": "bold"} in basic["content"][0]["marks"]
    assert {"type": "italic"} in basic["content"][1]["marks"]
    assert _style(_texts(rich)[-2])["color"] == "#0f766e"
    assert {"type": "underline"} in _texts(rich)[-1]["marks"]


@pytest.mark.parametrize(
    "attrs",
    [
        {"color": "url(file:///secret)"},
        {"color": "red;display:none"},
        {"fontSize": True},
        {"fontSize": float("nan")},
        {"fontSize": 999},
    ],
)
def test_typography_rejects_css_and_invalid_sizes(attrs):
    rich = reconcile_document({"skills": ["Python"]})
    rich["content"][-1]["content"][0]["content"][0]["marks"] = [
        {"type": "textStyle", "attrs": attrs}
    ]
    with pytest.raises(ValueError):
        validate_document(rich)


def test_revised_text_keeps_inline_typography_and_section_heading():
    rich = reconcile_document({"skills": ["Python SQL"]})
    section = rich["content"][-1]
    section["attrs"]["headingColor"] = "#17324d"
    section["content"][0]["content"][0]["marks"] = [
        {"type": "textStyle", "attrs": {"color": "#0f766e", "fontSize": 10}}
    ]
    before = deepcopy(rich)
    revised = reconcile_document({"skills": ["Python SQLite"]}, rich)
    assert revised["content"][-1]["attrs"]["headingColor"] == "#17324d"
    assert all(_style(n)["color"] == "#0f766e" for n in _texts(revised))
    assert document_content(revised)["skills"] == ("Python SQLite",)
    assert rich == before


def test_recovery_does_not_use_changed_original_or_overwrite_saved_formats(tmp_path):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    source = tmp_path / "source.pdf"
    _colored_pdf(source)
    profile = profiles.import_resume(
        workspace_id=workspace.workspace_id,
        source_path=source,
        mode=ResumeImportMode.MANAGED,
        index_original=True,
    )
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    editor = ResumeEditorService(db)
    plain = editor.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=base.content,
        document=reconcile_document(base.content),
    )
    assert plain.document == reconcile_document(base.content)
    managed = Path(
        profiles.load_document(
            workspace_id=workspace.workspace_id,
            document_id=profile.document_id,
        ).managed_path
    )
    managed.write_bytes(managed.read_bytes() + b"\nchanged\n")
    loaded = editor.get_version(
        workspace_id=workspace.workspace_id, version_id=base.version_id
    )
    assert loaded.document is None
    assert loaded.content == base.content
