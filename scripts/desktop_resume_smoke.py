"""Fictional fixture and independent export checks for packaged desktop CI."""

import sys
from pathlib import Path

from docx import Document
from docx.shared import Pt, RGBColor
from pdfminer.high_level import extract_pages
from pdfminer.layout import LTChar
from pypdf import PdfReader
from reportlab.pdfgen import canvas


def create(folder: Path) -> None:
    folder.mkdir(parents=True, exist_ok=True)
    pdf = canvas.Canvas(str(folder / "fictional.pdf"))
    pdf.setFillColorRGB(0.1, 0.2, 0.3)
    pdf.setFont("Helvetica-Bold", 24)
    pdf.drawString(50, 780, "Example Candidate")
    pdf.setFillColorRGB(0, 0.5, 0.4)
    pdf.setFont("Helvetica", 12)
    pdf.drawString(50, 750, "AI Engineer")
    pdf.setFont("Helvetica-Bold", 14)
    pdf.drawString(50, 710, "Skills")
    pdf.setFont("Helvetica", 10)
    pdf.setFillColorRGB(0.1, 0.3, 0.8)
    pdf.drawString(50, 690, "Python FastAPI")
    pdf.save()
    word = Document()
    run = word.add_paragraph().add_run("Word Candidate")
    run.bold = True
    run.font.size = Pt(24)
    run.font.color.rgb = RGBColor.from_string("17324D")
    word.add_paragraph("Skills")
    run = word.add_paragraph().add_run("Python SQL")
    run.font.color.rgb = RGBColor.from_string("0F766E")
    run.font.size = Pt(11)
    word.save(folder / "fictional.docx")


def verify(folder: Path, recovered: bool = False) -> None:
    prefix = "recovered" if recovered else "export"
    word = Document(folder / f"{prefix}.docx")
    text = "\n".join(p.text for p in word.paragraphs)
    assert "Packaged edit" in text, "Editor change missing from Word export"
    if recovered:
        assert "Recovered draft" in text, "Recovered draft missing from Word export"
    name = next(
        run
        for paragraph in word.paragraphs
        for run in paragraph.runs
        if run.text == "Example Candidate"
    )
    assert name.bold and name.font.size.pt == 24
    assert str(name.font.color.rgb) == "1A334C"
    reader = PdfReader(folder / f"{prefix}.pdf")
    assert len(reader.pages) == 1
    assert "Packaged edit" in reader.pages[0].extract_text()
    if recovered:
        assert "Recovered draft" in reader.pages[0].extract_text()
    chars = []

    def walk(node):
        if isinstance(node, LTChar):
            chars.append(node)
        elif hasattr(node, "__iter__"):
            for child in node:
                walk(child)

    for page in extract_pages(folder / f"{prefix}.pdf"):
        walk(page)
    assert any(round(char.size, 1) == 24 for char in chars)
    assert any(
        isinstance(char.graphicstate.ncolor, tuple)
        and len(char.graphicstate.ncolor) == 3
        and all(
            abs(actual - expected) < 1e-5
            for actual, expected in zip(
                char.graphicstate.ncolor, (26 / 255, 76 / 255, 204 / 255), strict=True
            )
        )
        for char in chars
    ), "Imported blue text missing from PDF export"
    print("Packaged Word/PDF exports retain edited text, colors and font sizes")


if __name__ == "__main__":
    if sys.argv[1] == "verify-draft":
        verify(Path(sys.argv[2]), recovered=True)
    else:
        {"create": create, "verify": verify}[sys.argv[1]](Path(sys.argv[2]))
