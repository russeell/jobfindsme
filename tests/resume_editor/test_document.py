from copy import deepcopy

import pytest
from docx import Document
from pypdf import PdfReader

from jobfindsme.resume_editor import ResumeEditorError, ResumeEditorService
from jobfindsme.resume_editor.document import (
    document_content,
    paragraph,
    reconcile_document,
    validate_document,
)
from tests.resume_editor.test_service import _confirmed_resume


def test_pdf_import_to_editable_word_preserves_original_and_formats(tmp_path):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    original = profiles.current_version(workspace_id=workspace.workspace_id)
    pdf = tmp_path / "source.pdf"
    service.export(
        workspace_id=workspace.workspace_id,
        version_id=original.version_id,
        destination=pdf,
        format="pdf",
        template="professional",
    )
    original_bytes = pdf.read_bytes()
    draft = profiles.import_resume(workspace_id=workspace.workspace_id, source_path=pdf)
    profiles.confirm_profile(
        workspace_id=workspace.workspace_id,
        profile_id=draft.profile_id,
        accepted_fact_ids=[fact.fact_id for fact in draft.facts],
    )
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    document = reconcile_document(base.content)
    skills = next(s for s in document["content"] if s["attrs"]["section"] == "skills")
    skills["content"] = [
        {
            "type": "bulletList",
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        {
                            "type": "paragraph",
                            "content": [
                                {
                                    "type": "text",
                                    "text": "Python",
                                    "marks": [{"type": "bold"}],
                                },
                                {"type": "hardBreak"},
                                {
                                    "type": "text",
                                    "text": "本机编辑验证",
                                    "marks": [{"type": "underline"}],
                                },
                            ],
                        }
                    ],
                }
            ],
        }
    ]
    edited = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=document_content(document),
        document=document,
    )
    loaded = ResumeEditorService(db).get_version(
        workspace_id=workspace.workspace_id, version_id=edited.version_id
    )
    assert loaded.document == validate_document(document)
    assert pdf.read_bytes() == original_bytes
    assert (
        service.get_version(
            workspace_id=workspace.workspace_id, version_id=base.version_id
        ).content
        == base.content
    )
    word = tmp_path / "edited.docx"
    rendered = tmp_path / "edited.pdf"
    for destination, format in ((word, "docx"), (rendered, "pdf")):
        service.export(
            workspace_id=workspace.workspace_id,
            version_id=edited.version_id,
            destination=destination,
            format=format,
            template="professional",
        )
    output = Document(word)
    assert any(
        run.bold and run.text == "Python" for p in output.paragraphs for run in p.runs
    )
    assert any(
        run.underline and run.text == "本机编辑验证"
        for p in output.paragraphs
        for run in p.runs
    )
    assert "本机编辑验证" in "".join(
        p.extract_text() for p in PdfReader(rendered).pages
    )
    assert output.sections[0].page_width.cm == pytest.approx(21, abs=0.01)


def test_legacy_edit_and_ai_changes_preserve_only_unchanged_entry_formatting(tmp_path):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    document = reconcile_document(base.content)
    projects = next(
        s for s in document["content"] if s["attrs"]["section"] == "projects"
    )
    projects["content"][0]["content"][0]["marks"] = [{"type": "bold"}]
    rich = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=base.version_id,
        content=base.content,
        document=document,
    )
    changed = service.save_edit(
        workspace_id=workspace.workspace_id,
        base_version_id=rich.version_id,
        content={**rich.content, "skills": ["编辑后的技能"]},
    )
    unchanged = next(
        s for s in changed.document["content"] if s["attrs"]["section"] == "projects"
    )
    assert unchanged == projects
    restored = service.restore_version(
        workspace_id=workspace.workspace_id, version_id=rich.version_id
    )
    assert restored.document == rich.document
    plain = service.restore_version(
        workspace_id=workspace.workspace_id, version_id=base.version_id
    )
    assert '"marks"' not in str(plain.document).replace("'", '"')


@pytest.mark.parametrize(
    "mutation",
    [
        lambda d: d["content"].pop(),
        lambda d: d["content"][0].update(attrs={"section": "unknown"}),
        lambda d: d["content"][0].update(
            content=[{"type": "image", "attrs": {"src": "file:///private"}}]
        ),
        lambda d: d["content"][0].update(
            content=[
                {
                    "type": "paragraph",
                    "content": [
                        {
                            "type": "text",
                            "text": "link",
                            "marks": [
                                {
                                    "type": "link",
                                    "attrs": {"href": "javascript:alert(1)"},
                                }
                            ],
                        }
                    ],
                }
            ]
        ),
    ],
)
def test_invalid_document_cannot_create_version(tmp_path, mutation):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    document = reconcile_document(base.content)
    mutation(document)
    with pytest.raises(ResumeEditorError, match="invalid resume document"):
        service.save_edit(
            workspace_id=workspace.workspace_id,
            base_version_id=base.version_id,
            content=base.content,
            document=document,
        )
    assert len(service.list_versions(workspace_id=workspace.workspace_id)) == 1


def test_document_plain_text_mismatch_is_rejected(tmp_path):
    db, workspace, profiles = _confirmed_resume(tmp_path)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    with pytest.raises(ResumeEditorError, match="does not match"):
        ResumeEditorService(db).save_edit(
            workspace_id=workspace.workspace_id,
            base_version_id=base.version_id,
            content={"skills": ["different"]},
            document=reconcile_document(base.content),
        )


def test_nested_lists_preserve_order_and_multiline():
    document = reconcile_document({"skills": ["one\ntwo"]})
    skills = document["content"][-1]
    skills["content"] = [
        {
            "type": "orderedList",
            "attrs": {"start": 3},
            "content": [
                {
                    "type": "listItem",
                    "content": [
                        paragraph("one"),
                        {
                            "type": "bulletList",
                            "content": [
                                {
                                    "type": "listItem",
                                    "content": [paragraph("two\nthree")],
                                }
                            ],
                        },
                    ],
                }
            ],
        }
    ]
    assert document_content(validate_document(deepcopy(document)))["skills"] == (
        "3. one\n• two\nthree",
    )


def test_legacy_private_bullets_load_and_export_without_rewriting_history(tmp_path):
    import json

    db, workspace, profiles = _confirmed_resume(tmp_path)
    service = ResumeEditorService(db)
    base = profiles.current_version(workspace_id=workspace.workspace_id)
    content = {key: [] for key in base.content}
    content["experience"] = ["  \uf0b7 完成检索服务\n\uf0b7 完成性能测试"]
    document = reconcile_document(content)
    block = next(
        s for s in document["content"] if s["attrs"]["section"] == "experience"
    )["content"][0]
    # Whitespace and bullet can belong to different formatted text runs.
    block["content"] = [
        {"type": "text", "text": "  "},
        {"type": "text", "text": "\uf0b7 完成检索服务", "marks": [{"type": "bold"}]},
        {"type": "hardBreak"},
        {"type": "text", "text": "\uf0b7 完成性能测试"},
    ]
    stored_content = json.dumps(content)
    stored_document = json.dumps(document)
    with db.connect() as conn:
        conn.execute(
            "UPDATE resume_versions SET content_json=?,document_json=? "
            "WHERE version_id=?",
            (stored_content, stored_document, base.version_id),
        )
    loaded = service.get_version(
        workspace_id=workspace.workspace_id, version_id=base.version_id
    )
    assert loaded.content["experience"] == ("  • 完成检索服务\n• 完成性能测试",)
    assert (
        profiles.current_version(workspace_id=workspace.workspace_id).content
        == loaded.content
    )
    rich = next(
        s for s in loaded.document["content"] if s["attrs"]["section"] == "experience"
    )["content"][0]
    assert rich["content"][1]["marks"] == [{"type": "bold"}]
    assert rich["content"][1]["text"] == "• 完成检索服务"
    assert rich["content"][3]["text"] == "• 完成性能测试"
    assert "\uf0b7" in document_content(document)["experience"][0]
    for format in ("docx", "pdf"):
        output = tmp_path / f"portable.{format}"
        service.export(
            workspace_id=workspace.workspace_id,
            version_id=base.version_id,
            destination=output,
            format=format,
            template="professional",
        )
        text = (
            "\n".join(p.text for p in Document(output).paragraphs)
            if format == "docx"
            else "\n".join(p.extract_text() for p in PdfReader(output).pages)
        )
        assert "• 完成检索服务" in text
        assert "• 完成性能测试" in text
        assert "\uf0b7" not in text
    with db.connect() as conn:
        row = conn.execute(
            "SELECT content_json,document_json FROM resume_versions WHERE version_id=?",
            (base.version_id,),
        ).fetchone()
    assert tuple(row) == (stored_content, stored_document)
